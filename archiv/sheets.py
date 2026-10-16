"""Google Sheets: the whole spreadsheet as .xlsx, kept as a file and as rows in the private schema.

A sheet shared as "anyone with the link can view" exports without a login. Each tab becomes rows of
{column header: value}: the header row is the first row with at least two text cells, columns without a
header are named by their letter. Values are the ones the spreadsheet computed (dates as ISO text).
"""

import hashlib
import io
import json
import re
import urllib.request
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ID = re.compile(r"/spreadsheets/d/([A-Za-z0-9_-]+)")


class SheetError(Exception):
    pass


def export_url(url: str) -> str:
    """The .xlsx export of a sheet link (edit, view or share link); anything else is used as it is."""
    match = ID.search(url)
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=xlsx" if match else url


def download(url: str) -> bytes:
    with urllib.request.urlopen(export_url(url), timeout=60) as response:
        data = response.read()
    # Not shared: Google answers with its sign-in page instead of a file
    if not data.startswith(b"PK"):
        raise SheetError("no spreadsheet came back - is the sheet shared as 'anyone with the link'?")
    return data


def value(cell):
    if cell is None:
        return None
    if isinstance(cell, datetime):
        return cell.date().isoformat() if cell.time() == time(0) else cell.isoformat()
    if isinstance(cell, (date, time)):
        return cell.isoformat()
    if isinstance(cell, str):
        cell = cell.strip()
        return cell or None
    return cell


def read(data: bytes) -> dict[str, list[tuple[int, dict]]]:
    """Tab name -> [(row number in the sheet, {header: value})]."""
    book = load_workbook(io.BytesIO(data), data_only=True)
    tabs = {}
    for sheet in book.worksheets:
        rows = [(n, [value(c) for c in cells]) for n, cells in enumerate(sheet.iter_rows(values_only=True), start=1)]
        header_at = next((i for i, (_, cells) in enumerate(rows)
                          if sum(isinstance(c, str) for c in cells) >= 2), None)
        if header_at is None:
            tabs[sheet.title] = []
            continue
        keys, seen = [], {}
        for col, head in enumerate(rows[header_at][1], start=1):
            key = str(head) if head is not None else get_column_letter(col)
            seen[key] = seen.get(key, 0) + 1
            keys.append(key if seen[key] == 1 else f"{key}_{seen[key]}")
        tabs[sheet.title] = [
            (n, {k: v for k, v in zip(keys, cells) if v is not None})
            for n, cells in rows[header_at + 1:]
            if any(v is not None for v in cells)
        ]
    return tabs


def fingerprint(tabs: dict) -> str:
    """Changes only when a value changes, unlike the file Google builds fresh on every export."""
    return hashlib.sha256(json.dumps(tabs, sort_keys=True, default=str).encode()).hexdigest()
